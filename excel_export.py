from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


def generate_students_excel(students):

    columns = [
        "Student ID",
        "Student Name",
        "Email",
        "Phone",
        "Course",
        "Semester",
        "Section",
        "Gender"
    ]

    dataframe = pd.DataFrame(
        students,
        columns=columns
    )

    excel_buffer = BytesIO()

    with pd.ExcelWriter(
        excel_buffer,
        engine="openpyxl"
    ) as writer:

        dataframe.to_excel(
            writer,
            sheet_name="Students",
            index=False
        )

        worksheet = writer.sheets["Students"]

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        for cell in worksheet[1]:

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        column_widths = {
            "A": 15,
            "B": 24,
            "C": 30,
            "D": 16,
            "E": 22,
            "F": 12,
            "G": 12,
            "H": 14
        }

        for column, width in column_widths.items():
            worksheet.column_dimensions[column].width = width

        for row in worksheet.iter_rows(
            min_row=2
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center"
                )

    excel_buffer.seek(0)

    return excel_buffer

def generate_attendance_excel(attendance_records):

    columns = [
        "Record ID",
        "Student ID",
        "Student Name",
        "Attendance Date",
        "Status",
    ] 

    dataframe = pd.DataFrame(
        attendance_records,
        columns=columns
    )

    excel_buffer = BytesIO()

    with pd.ExcelWriter(
        excel_buffer,
        engine="openpyxl"
    ) as writer:

        dataframe.to_excel(
            writer,
            sheet_name="Attendance",
            index=False
        )

        worksheet = writer.sheets["Attendance"]

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="16A34A"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        column_widths = {
            "A": 12,
            "B": 15,
            "C": 24,
            "D": 18,
            "E": 14
        }

        for column, width in column_widths.items():
            worksheet.column_dimensions[column].width = width

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center"
                )

    excel_buffer.seek(0)

    return excel_buffer

def generate_marks_excel(marks_records):

    columns = [
        "Record ID",
        "Student ID",
        "Student Name",
        "Subject",
        "Internal Marks",
        "External Marks",
        "Total Marks",
        "Percentage",
        "Grade",
        "Created At"
    ]

    dataframe = pd.DataFrame(
        marks_records,
        columns=columns
    )

    excel_buffer = BytesIO()

    with pd.ExcelWriter(
        excel_buffer,
        engine="openpyxl"
    ) as writer:

        dataframe.to_excel(
            writer,
            sheet_name="Marks",
            index=False
        )

        worksheet = writer.sheets["Marks"]

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="7C3AED"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        column_widths = {
            "A": 12,
            "B": 15,
            "C": 24,
            "D": 22,
            "E": 16,
            "F": 16,
            "G": 14,
            "H": 14,
            "I": 12,
            "J": 22
        }

        for column, width in column_widths.items():
            worksheet.column_dimensions[column].width = width

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center"
                )

    excel_buffer.seek(0)

    return excel_buffer